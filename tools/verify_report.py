"""Verify the filled report: structure, content and figures."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

REPORT = Path("docs/EIE4428-mini-project-report-RVM.docx")

document = Document(str(REPORT))

print(f"file: {REPORT}  ({REPORT.stat().st_size} bytes)")
print(f"paragraphs: {len(document.paragraphs)}   tables: {len(document.tables)}")
print(f"sections: {len(document.sections)}")

images = document.element.body.findall(f".//{qn('w:drawing')}")
print(f"inline images: {len(images)}")

words = sum(len(p.text.split()) for p in document.paragraphs)
print(f"approx words: {words}")

print()
print("=== TABLE 0 (group members) ===")
for row in document.tables[0].rows:
    print("  ", [c.text for c in row.cells])

print()
print("=== TABLE 1 (functionalities) ===")
for row in document.tables[1].rows:
    left = row.cells[0].text.strip()
    right = row.cells[1].text.strip()
    print(f"  {left[:34]:36s} | {right[:88]}")

print()
print("=== HEADINGS / ANCHORS ===")
markers = (
    "Summary of the Project",
    "Details of the Project",
    "Achievements of the Project",
    "Methodology",
    "Implementation Details",
    "Design of User Interface",
    "Co-operation between Group Members",
    "Conclusions",
    "References",
)
for marker in markers:
    hits = [i for i, p in enumerate(document.paragraphs) if p.text.strip().startswith(marker)]
    print(f"  {marker:38s} -> {hits}")

print()
print("=== PLACEHOLDER CHECK (instructional text that should be gone) ===")
leftovers = []
for index, paragraph in enumerate(document.paragraphs):
    text = paragraph.text.strip()
    if text.startswith("(Describe") or text.startswith("(Give") or text.startswith("(Make") or "(video only, or audio" in text:
        leftovers.append((index, text[:70]))
print("  leftovers:", leftovers if leftovers else "none")

print()
print("=== FIRST PARAGRAPHS OF EACH MAJOR SECTION ===")
for index, paragraph in enumerate(document.paragraphs):
    text = paragraph.text.strip()
    if any(text.startswith(m) for m in ("Summary of the Project", "Methodology", "Conclusions", "References")):
        print(f"  [{index}] {text[:110]}")