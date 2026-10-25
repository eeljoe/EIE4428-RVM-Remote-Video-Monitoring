"""Inspect the report template: paragraphs, styles and tables."""
from pathlib import Path

from docx import Document

TEMPLATE = Path("EIE4428-mini-project-report-template.docx")

doc = Document(str(TEMPLATE))

print("=== PARAGRAPHS ===")
for index, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip()
    print(f"[{index:3d}] style={paragraph.style.name!r:22s} | {text[:95]}")

print()
print("=== TABLES ===")
for t_index, table in enumerate(doc.tables):
    print(f"--- table {t_index}: {len(table.rows)} rows x {len(table.columns)} cols")
    for r_index, row in enumerate(table.rows):
        cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
        print(f"    row {r_index}: {cells}")